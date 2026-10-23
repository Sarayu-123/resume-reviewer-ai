"""Robust text extraction for DOCX resumes, including tables and headers."""

from __future__ import annotations

import logging
import re
import unicodedata
from pathlib import Path
from typing import BinaryIO, Union

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

logger = logging.getLogger(__name__)
DocxInput = Union[str, Path, BinaryIO]


def _clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text or "")
    text = text.replace("\u00a0", " ").replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _iter_block_items(parent):
    """Yield paragraphs and tables in their original document order."""
    if hasattr(parent, "iter_inner_content"):
        yield from parent.iter_inner_content()
        return

    # Compatibility fallback for older python-docx versions.
    from docx.oxml.text.paragraph import CT_P
    from docx.oxml.table import CT_Tbl

    from docx.text.paragraph import Paragraph as ParagraphClass
    from docx.table import Table as TableClass

    for child in parent.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield ParagraphClass(child, parent)
        elif isinstance(child, CT_Tbl):
            yield TableClass(child, parent)


def _table_to_text(table: Table) -> list[str]:
    rows = []
    for row in table.rows:
        cells = []
        seen = set()
        for cell in row.cells:
            value = _clean_text(" ".join(p.text for p in cell.paragraphs))
            # Merged Word cells can appear more than once in row.cells.
            if value and value not in seen:
                cells.append(value)
                seen.add(value)
        if cells:
            rows.append(" | ".join(cells))
    return rows


def read_docx(file: DocxInput) -> str:
    """Extract paragraphs, tables, and header/footer text from a DOCX file."""
    try:
        document = Document(file)
    except Exception as exc:
        logger.exception("Failed to open DOCX resume")
        raise ValueError("Could not read this DOCX. Check that it is a valid Word document.") from exc

    output: list[str] = []

    try:
        for block in _iter_block_items(document):
            if isinstance(block, Paragraph):
                value = _clean_text(block.text)
                if value:
                    output.append(value)
            elif isinstance(block, Table):
                output.extend(_table_to_text(block))

        # Headers and footers often contain name/contact details. Add unique
        # non-empty text without repeating the same header on every page.
        seen_header_footer = set()
        for section in document.sections:
            for part in (section.header, section.footer):
                for paragraph in part.paragraphs:
                    value = _clean_text(paragraph.text)
                    if value and value not in seen_header_footer:
                        output.append(value)
                        seen_header_footer.add(value)
                for table in part.tables:
                    for value in _table_to_text(table):
                        if value not in seen_header_footer:
                            output.append(value)
                            seen_header_footer.add(value)

    except Exception as exc:
        logger.exception("Failed while extracting DOCX content")
        raise ValueError("The DOCX opened, but its text could not be extracted.") from exc

    # Preserve paragraph/row boundaries for section and skill detection.
    result = "\n".join(output)
    result = re.sub(r"\n{3,}", "\n\n", result).strip()
    return result if len(re.sub(r"\s", "", result)) >= 10 else ""
